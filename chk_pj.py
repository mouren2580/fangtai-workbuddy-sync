# -*- coding: utf-8 -*-
"""确认 CSM配件 是否与 CSM服务项目 内容相同，以及销售分类取值分布。"""
from openpyxl import load_workbook
from collections import Counter
EXCEL = r"D:\WorkBuddy\2026年10月西北服务产品.xlsx"
wb = load_workbook(EXCEL, read_only=True, data_only=True)
def rows(nm):
    return [r for r in wb[nm].iter_rows(values_only=True)]
a = rows("CSM配件")[2:]
b = rows("CSM服务项目")[2:]
print("CSM配件 行数=%d  CSM服务项目 行数=%d" % (len(a), len(b)))
same = sum(1 for x, y in zip(a, b) if x == y)
print("逐行完全相同的前 %d 行中: %d 行一致" % (min(len(a), len(b)), same))
print("\n前3行对比:")
for i in range(3):
    print("  配件: %s" % (str(a[i][:8])))
    print("  项目: %s" % (str(b[i][:8])))
print("\nCSM配件 col16(销售分类) 取值 TOP15:", Counter(str(r[16]) for r in a).most_common(15))
print("CSM配件 col13(服务收费项目) 取值 TOP10:", Counter(str(r[13]) for r in a).most_common(10))
print("\nCSM服务项目 col16(销售分类) TOP15:", Counter(str(r[16]) for r in b).most_common(15))
wb.close()
