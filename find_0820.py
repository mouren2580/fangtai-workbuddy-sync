import openpyxl, datetime

f = r"D:\WorkBuddy\2026年8月西北服务产品.xlsx"
wb = openpyxl.load_workbook(f, data_only=True)

def serial_to_date(s):
    return datetime.datetime(1899, 12, 30) + datetime.timedelta(days=s)

target = datetime.datetime(2026, 8, 20)
print("扫描所有 sheet 中 2026-08-20 出现位置...")

for ws in wb.worksheets:
    rows = list(ws.iter_rows(values_only=True))
    hdr = None
    for i, row in enumerate(rows[:6]):
        if row and any(isinstance(c, str) and ('日期' in c or '时间' in c) for c in row):
            hdr = row
            break
    found = []
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            if v is None:
                continue
            d = None
            if isinstance(v, datetime.datetime):
                d = v
            elif isinstance(v, (int, float)):
                if 40000 < v < 70000:
                    d = serial_to_date(v)
            if d and d.year == 2026 and d.month == 8 and d.day == 20:
                col_name = hdr[j] if hdr and j < len(hdr) else f"列{j+1}"
                found.append((i + 1, col_name, v))
    if found:
        print(f"\n[{ws.title}] 发现 2026-08-20 共 {len(found)} 次，前10条：")
        for r, c, v in found[:10]:
            print(f"  行{r} {c} = {v}")

# 同时重新输出各表日期列最大值（含列名）
print("\n\n各表日期/时间列最大值：")
for ws in wb.worksheets:
    rows = list(ws.iter_rows(values_only=True))
    hdr = None
    date_cols = []
    sheet_max = None
    max_col = None
    for i, row in enumerate(rows):
        if hdr is None:
            if i < 6 and row and any(isinstance(c, str) and ('日期' in c or '时间' in c) for c in row):
                hdr = row
                date_cols = [(j, c) for j, c in enumerate(row) if isinstance(c, str) and ('日期' in c or '时间' in c)]
                continue
            if i >= 6 and hdr is None:
                break
        if hdr is not None:
            for j, col_name in date_cols:
                if j < len(row):
                    v = row[j]
                    if v is None:
                        continue
                    d = None
                    if isinstance(v, datetime.datetime):
                        d = v
                    elif isinstance(v, (int, float)) and 40000 < v < 70000:
                        d = serial_to_date(v)
                    if d:
                        d = d.replace(hour=0, minute=0, second=0, microsecond=0)
                        if sheet_max is None or d > sheet_max:
                            sheet_max = d
                            max_col = col_name
    if sheet_max:
        print(f"  [{ws.title}] {max_col} max={sheet_max.date()}")
