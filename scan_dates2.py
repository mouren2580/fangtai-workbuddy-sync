# -*- coding: utf-8 -*-
"""扫 Excel 各表的最大日期与总行数，判断各表真实截止日是否一致。"""
import sys, datetime
from openpyxl import load_workbook

EXCEL = r"D:\WorkBuddy\2026年10月西北服务产品.xlsx"
TARGETS = ["工单", "CSM服务项目", "CSM配件", "工单自购配件明细", "服务产品收入统计", "WMS买断"]

def is_date(v):
    return isinstance(v, (datetime.datetime, datetime.date)) and not isinstance(v, bool)

def pick_date_cols(header):
    """表头里含 日期/时间 的列索引（0基）"""
    out = []
    for i, h in enumerate(header):
        h = str(h or "")
        if ("日期" in h) or ("时间" in h):
            out.append(i)
    return out

wb = load_workbook(EXCEL, read_only=True, data_only=True)
print("工作表:", wb.sheetnames)
for name in wb.sheetnames:
    if name not in TARGETS:
        continue
    ws = wb[name]
    it = ws.iter_rows(values_only=True)
    try:
        header = next(it)
    except StopIteration:
        print(f"\n[{name}] 空表"); continue
    head = [str(h or "") for h in header]
    dcols = pick_date_cols(head)
    if not dcols:
        print(f"\n[{name}] 未找到日期列，表头前20: {head[:20]}"); continue
    rows = 0
    mx = {}
    for r in it:
        if r is None or all(c is None for c in r):
            continue
        rows += 1
        for c in dcols:
            if c < len(r) and is_date(r[c]):
                d = r[c]
                d = d.date() if isinstance(d, datetime.datetime) else d
                if c not in mx or d > mx[c]:
                    mx[c] = d
    print(f"\n[{name}] 有效行数={rows}")
    for c in dcols:
        print(f"   col{c} {head[c]!r}: max={mx.get(c)}")
wb.close()
