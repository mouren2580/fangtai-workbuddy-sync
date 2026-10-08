# -*- coding: utf-8 -*-
import datetime
from openpyxl import load_workbook

EXCEL = r"D:\WorkBuddy\2026年10月西北服务产品.xlsx"
wb = load_workbook(EXCEL, read_only=True, data_only=True)

def is_date(v):
    return isinstance(v, (datetime.datetime, datetime.date)) and not isinstance(v, bool)

def d(v):
    return v.date() if isinstance(v, datetime.datetime) else v

for name in ["CSM服务项目", "CSM配件", "服务产品收入统计"]:
    ws = wb[name]
    print(f"\n===== {name} (dims={ws.calculate_dimension()}) =====")
    for k, r in enumerate(ws.iter_rows(values_only=True)):
        if k >= 4: break
        cells = [str(c)[:14] if c is not None else "" for c in r[:14]]
        print(f"  row{k}: {cells}")
    # 找表头行：第一个非空单元格数 >= 8 的行
    ws2 = wb[name]
    hdr = None; hr = None
    for k, r in enumerate(ws2.iter_rows(values_only=True)):
        n = sum(1 for c in r if c is not None and str(c).strip())
        if n >= 8:
            hdr, hr = r, k; break
    if hdr is None:
        print("  未定位到表头"); continue
    head = [str(c or "") for c in hdr]
    dcols = [i for i, h in enumerate(head) if "日期" in h or "时间" in h]
    print(f"  表头行={hr}, 日期列: {[(i, head[i]) for i in dcols]}")
    ws3 = wb[name]
    rows = 0; mx = {}
    for k, r in enumerate(ws3.iter_rows(values_only=True)):
        if k <= hr: continue
        if r is None or all(c is None for c in r): continue
        rows += 1
        for i in dcols:
            if i < len(r) and is_date(r[i]):
                v = d(r[i])
                if i not in mx or v > mx[i]: mx[i] = v
    print(f"  有效行数={rows}")
    for i in dcols:
        print(f"    col{i} {head[i]!r}: max={mx.get(i)}")
wb.close()
