import openpyxl, datetime, re

f = r"D:\WorkBuddy\2026年8月西北服务产品.xlsx"
wb = openpyxl.load_workbook(f, data_only=True)

def serial_to_date(s):
    return datetime.datetime(1899, 12, 30) + datetime.timedelta(days=s)

target = datetime.datetime(2026, 8, 20).date()
patterns = [
    re.compile(r'2026\s*[-/年]?\s*0?8\s*[-/月]?\s*0?20\s*日?'),
    re.compile(r'0?8\s*[-/月]?\s*0?20\s*[-/日]?\s*2026'),
    re.compile(r'20260820'),
]

def looks_like_date_string(s):
    for p in patterns:
        if p.search(s):
            return True
    return False

print("扫描所有 sheet 中可能被解析为 2026-08-20 的单元格...")
any_found = False
for ws in wb.worksheets:
    rows = list(ws.iter_rows(values_only=True))
    found = []
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            if v is None:
                continue
            d = None
            reason = ""
            if isinstance(v, datetime.datetime):
                if v.date() == target:
                    d = v
                    reason = "datetime"
            elif isinstance(v, (int, float)):
                if 40000 < v < 70000:
                    dt = serial_to_date(v)
                    if dt.date() == target:
                        d = dt
                        reason = f"serial({v})"
            elif isinstance(v, str):
                s = v.strip()
                if looks_like_date_string(s):
                    d = s
                    reason = "string"
            if d:
                col_name = ""
                for hi in range(min(len(rows), 6)):
                    if rows[hi] and j < len(rows[hi]):
                        c = rows[hi][j]
                        if isinstance(c, str):
                            col_name = c
                            break
                found.append((i + 1, j + 1, col_name, reason, v))
    if found:
        any_found = True
        print(f"\n[{ws.title}] 发现 {len(found)} 个可疑单元格：")
        for r, c, col, reason, v in found:
            print(f"  行{r} 列{c}({col}) [{reason}] = {v!r}")

if not any_found:
    print("\n未发现任何可能被解析为 2026-08-20 的单元格。")

# 同时输出真正的业务日期最大值（仅日期格式/serial）
print("\n\n各表真实日期列最大值（排除字符串）：")
for ws in wb.worksheets:
    rows = list(ws.iter_rows(values_only=True))
    hdr = None
    date_cols = []
    sheet_max = None
    max_col = None
    for i, row in enumerate(rows):
        if hdr is None:
            if i < 6 and row and any(isinstance(c, str) and ('日期' in c or '时间' in c) for c in row):
                hdr = row
                date_cols = [(j, c) for j, c in enumerate(row) if isinstance(c, str) and ('日期' in c or '时间' in c)]
                continue
            if i >= 6 and hdr is None:
                break
        if hdr is not None:
            for j, col_name in date_cols:
                if j < len(row):
                    v = row[j]
                    if v is None:
                        continue
                    d = None
                    if isinstance(v, datetime.datetime):
                        d = v
                    elif isinstance(v, (int, float)) and 40000 < v < 70000:
                        d = serial_to_date(v)
                    if d:
                        d = d.replace(hour=0, minute=0, second=0, microsecond=0)
                        if sheet_max is None or d > sheet_max:
                            sheet_max = d
                            max_col = col_name
    if sheet_max:
        print(f"  [{ws.title}] {max_col} max={sheet_max.date()}")
