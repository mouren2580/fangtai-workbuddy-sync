# -*- coding: utf-8 -*-
"""打印各表完整表头 + 维度，用于核对列结构是否变化。"""
from openpyxl import load_workbook
EXCEL = r"D:\WorkBuddy\2026年10月西北服务产品.xlsx"
wb = load_workbook(EXCEL, read_only=True, data_only=True)
for name in wb.sheetnames:
    ws = wb[name]
    print("\n===== %s  dims=%s =====" % (name, ws.calculate_dimension()))
    rows = []
    for k, r in enumerate(ws.iter_rows(values_only=True)):
        rows.append(r)
        if k >= 2:
            break
    hdr = rows[1] if len(rows) > 1 else []
    for i, h in enumerate(hdr or []):
        print("  %3d %s" % (i, h))
wb.close()
