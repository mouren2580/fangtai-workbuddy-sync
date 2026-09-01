# -*- coding: utf-8 -*-
import openpyxl, datetime, sys, io

XLSX = r"D:\WorkBuddy\2026年9月西北服务产品.xlsx"
out = io.StringIO()
def log(*a):
    print(*a, file=out)

wb = openpyxl.load_workbook(XLSX, data_only=True)
log("=== 工作表清单 ===")
log("共 %d 个表: %s" % (len(wb.worksheets), [ws.title for ws in wb.worksheets]))

def serial_to_date(s):
    try:
        return datetime.datetime(1899,12,30)+datetime.timedelta(days=int(s))
    except Exception:
        return None

overall_max = None
overall_min = None
log("\n=== 各表日期列(含'日期'/'时间')扫描（严格只认 datetime 类型）===")
for ws in wb.worksheets:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        log("[%s] 空表" % ws.title); continue
    hdr_idx = -1
    for i, row in enumerate(rows[:8]):
        if row and any(('日期' in str(c) or '时间' in str(c)) for c in row if c is not None):
            hdr_idx = i; break
    if hdr_idx < 0:
        log("[%s] 无日期/时间列" % ws.title); continue
    hdr = rows[hdr_idx]
    cols = [(j, str(c)) for j, c in enumerate(hdr) if c and ('日期' in str(c) or '时间' in str(c))]
    dmin = None; dmax = None; hit = 0
    for row in rows[hdr_idx+1:]:
        if not row: continue
        for j, _ in cols:
            v = row[j] if j < len(row) else None
            if v is None or v == '': continue
            d = None
            if isinstance(v, datetime.datetime): d = v.date()
            elif isinstance(v, datetime.date): d = v
            if d is None: continue
            # 过滤明显异常年份（如 2031）
            if d.year < 2026 or d.year > 2027:
                continue
            hit += 1
            if dmin is None or d < dmin: dmin = d
            if dmax is None or d > dmax: dmax = d
    log("[%s] 日期列=%s  命中%d个  范围 %s ~ %s" % (ws.title, [c[1] for c in cols], hit, dmin, dmax))
    if dmax:
        if overall_max is None or dmax > overall_max: overall_max = dmax
        if overall_min is None or dmin < overall_min: overall_min = dmin

log("\n=== 整体业务日期范围（过滤异常年后）===")
log("MIN =", overall_min, " MAX =", overall_max)

# 专项：服务产品收入统计 表的日期列（看板核心）
log("\n=== 专项：服务产品收入统计 表 ===")
ws2 = wb["服务产品收入统计"] if "服务产品收入统计" in wb.sheetnames else None
if ws2 is None:
    log("未找到『服务产品收入统计』表！现有表名: %s" % [w.title for w in wb.worksheets])
else:
    rows = list(ws2.iter_rows(values_only=True))
    hdr_idx = -1
    for i, row in enumerate(rows[:8]):
        if row and any(('日期' in str(c) or '时间' in str(c)) for c in row if c is not None):
            hdr_idx = i; break
    log("表头行索引=%d, 表行数=%d" % (hdr_idx, len(rows)))
    if hdr_idx >= 0:
        hdr = rows[hdr_idx]
        cols = [(j, str(c)) for j, c in enumerate(hdr) if c and ('日期' in str(c) or '时间' in str(c))]
        dmax = None
        for row in rows[hdr_idx+1:]:
            if not row: continue
            for j, _ in cols:
                v = row[j] if j < len(row) else None
                if v is None or v == '': continue
                d = None
                if isinstance(v, datetime.datetime): d = v.date()
                elif isinstance(v, datetime.date): d = v
                if d and 2026 <= d.year <= 2027 and (dmax is None or d > dmax):
                    dmax = d
        log("该表最大业务日期 =", dmax)

# 输出文件
with open(r"D:\WorkBuddy\scan_sep_result.txt", "w", encoding="utf-8") as f:
    f.write(out.getvalue())
print(out.getvalue())
